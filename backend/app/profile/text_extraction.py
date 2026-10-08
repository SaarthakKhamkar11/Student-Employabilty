"""Resume text extraction: plain text, .txt, .pdf (pdfplumber) or .docx (python-docx).

Malformed input is rejected with a clear reason (ResumeRejected) rather than
passed on as garbage: empty text, binary content, text that is neither UTF-8
nor BOM-marked UTF-16, unsupported formats, files over 5 MB, and PDFs without
a text layer (scanned PDFs need OCR, which is not supported).
"""

import base64
import binascii
import io
import unicodedata
from pathlib import PurePath
from typing import Literal

# Imported at module load (not on first use): their first import takes about 1 s, which would
# otherwise land on the first request and could exceed a tool timeout.
import docx
import pdfplumber

MAX_FILE_BYTES = 5 * 1024 * 1024
MIN_PRINTABLE_RATIO = 0.85
SourceFormat = Literal["text", "txt", "pdf", "docx"]


class ResumeRejected(ValueError):
    """The resume cannot be read. Deterministic, so retrying cannot help."""


def _check_text(text: str) -> str:
    text = unicodedata.normalize("NFC", text).replace("\x00", "")
    if not text.strip():
        raise ResumeRejected("resume is empty or contains only whitespace")
    visible = [c for c in text if not c.isspace()]
    printable = sum(1 for c in visible if c.isprintable() and unicodedata.category(c) not in ("Co", "Cn", "Cs"))
    if printable / len(visible) < MIN_PRINTABLE_RATIO:
        raise ResumeRejected("resume looks like binary or garbled data, not text")
    return text


def _decode(data: bytes) -> str:
    if b"\x00" in data[:1024] and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
        raise ResumeRejected("file contains binary data, not text")
    for encoding in ("utf-8-sig", "utf-16"):
        if encoding == "utf-16" and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
            continue
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ResumeRejected("text file is not valid UTF-8 (or BOM-marked UTF-16); re-save it as UTF-8")


def _pdf_text(data: bytes) -> str:
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    except Exception as exc:
        raise ResumeRejected(f"PDF could not be read ({type(exc).__name__})") from exc
    if not text.strip():
        raise ResumeRejected("PDF has no text layer (scanned image?); OCR is not supported")
    return text


def _docx_text(data: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ResumeRejected(f"DOCX could not be read ({type(exc).__name__})") from exc
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        lines += [" | ".join(cell.text for cell in row.cells) for row in table.rows]
    return "\n".join(lines)


def extract_text(text: str | None = None, content_b64: str | None = None,
                 filename: str | None = None) -> tuple[str, SourceFormat]:
    if text is not None:
        return _check_text(text), "text"
    if content_b64 is None:
        raise ResumeRejected("no resume text or file was provided")
    try:
        data = base64.b64decode(content_b64, validate=True)
    except (binascii.Error, ValueError) as exc:
        raise ResumeRejected("file content is not valid base64") from exc
    if len(data) > MAX_FILE_BYTES:
        raise ResumeRejected(f"file is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB")
    suffix = PurePath(filename or "").suffix.lower()
    if suffix == ".pdf" or data.startswith(b"%PDF"):
        return _check_text(_pdf_text(data)), "pdf"
    if suffix == ".docx":
        return _check_text(_docx_text(data)), "docx"
    if suffix in ("", ".txt", ".text"):
        return _check_text(_decode(data)), "txt"
    raise ResumeRejected(f"unsupported file type '{suffix}' (use .pdf, .docx or .txt)")
