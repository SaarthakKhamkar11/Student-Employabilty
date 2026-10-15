"""Profile pipeline components: text extraction, the existing parser (and its .NET fix), skills, TF-IDF."""

import base64
import csv
import io

import pytest

from app.profile import sections
from app.profile.skills import ALIAS_TABLE_PATH, normalize_skills, skill_key
from app.profile.text_extraction import ResumeRejected, extract_text


def b64(data: bytes) -> str:
    return base64.b64encode(data).decode()


def minimal_pdf(text: str) -> bytes:
    """A real one-page PDF with a text layer (Helvetica), built by hand with a correct xref table."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
               b"/Resources << /Font << /F1 5 0 R >> >> >>",
               b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = b"%PDF-1.4\n", []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    return out + b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)


# -- text extraction ----------------------------------------------------------------------------
def test_text_txt_utf16_docx_and_pdf_are_extracted():
    assert extract_text(text="Skills: Python") == ("Skills: Python", "text")
    assert extract_text(content_b64=b64("Skills: Python, café".encode("utf-8-sig")), filename="cv.txt")[0] \
        == "Skills: Python, café"
    assert extract_text(content_b64=b64("Skills: SQL".encode("utf-16")), filename="cv.txt")[0] == "Skills: SQL"

    import docx

    document = docx.Document()
    document.add_paragraph("Skill Details")
    document.add_paragraph("Python, Power BI")
    buffer = io.BytesIO()
    document.save(buffer)
    text, fmt = extract_text(content_b64=b64(buffer.getvalue()), filename="cv.docx")
    assert fmt == "docx" and "Python, Power BI" in text

    text, fmt = extract_text(content_b64=b64(minimal_pdf("Skill Details Python SQL")), filename="cv.pdf")
    assert fmt == "pdf" and "Python SQL" in text


@pytest.mark.parametrize("kwargs", [
    {"text": ""}, {"text": "\x00\x01\x02\x03\x04 x"},
    {"content_b64": b64(b"\xff\xfe" + "\U0010ffff\U0010fffe".encode("utf-16-le")), "filename": "a.txt"},
    {"content_b64": b64(b"PK\x03\x04 not a docx"), "filename": "a.docx"},
    {"content_b64": b64(b"x" * (5 * 1024 * 1024 + 1)), "filename": "big.txt"},
    {},
])
def test_unusable_input_is_rejected_with_a_reason(kwargs):
    with pytest.raises(ResumeRejected) as info:
        extract_text(**kwargs)
    assert str(info.value)


# -- the existing parser is reused, and its .NET bug is fixed narrowly -----------------------------
def test_dotnet_keeps_its_leading_dot_and_other_trimming_is_unchanged():
    from src.parsing.resume_text_parser import extract_skills

    text = "Skill Details\n.NET- Exprience - 6 months\nPython.\n- SQL\n* Excel\n. Java\nCompany Details\nX"
    assert extract_skills(text) == [".NET- Exprience - 6 months", "Python", "SQL", "Excel", "Java"]


def test_section_extraction_delegates_to_the_existing_parser(monkeypatch):
    called = {}

    def fake_parse_resume(text):
        called["text"] = text
        return {"skills": [], "education_raw": "from src.parsing", "experience_raw": ""}

    monkeypatch.setattr(sections, "parse_resume", fake_parse_resume)
    result = sections.resume_sections("anything")
    assert called["text"] == "anything" and result["education"] == "from src.parsing"


def test_resume_sections_extract_projects_and_certifications_with_the_existing_slicer():
    result = sections.resume_sections("Projects\nChatbot\nSales dashboard\nCertifications\nAWS CCP\nDeclaration\nx")
    assert result["projects"] == ["Chatbot", "Sales dashboard"] and result["certifications"] == ["AWS CCP"]
    assert sections.resume_sections("Skills\nPython")["projects"] == "Not detected"


# -- skills ------------------------------------------------------------------------------------------
def test_alias_table_is_complete_unique_and_evidenced():
    with ALIAS_TABLE_PATH.open(newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    keys = [skill_key(r["alias"]) for r in rows]
    assert len(keys) == len(set(keys))
    assert {r["category"] for r in rows} <= {"technical", "soft"}
    assert all(r["evidence"] for r in rows)
    assert sum("students_unified.csv" in r["evidence"] or "jobs.csv" in r["evidence"] for r in rows) >= 200
    assert all(any(r["alias"] == r["canonical"] for r in rows if r["canonical"] == c)
               for c in {r["canonical"] for r in rows})  # every canonical name is itself a row


@pytest.mark.parametrize("raw, canonical, method", [
    ("JS", "JavaScript", "alias"), ("PowerBI", "Power BI", "alias"), ("React.js", "React", "alias"),
    ("ML", "Machine Learning", "alias"), ("JAVA", "Java", "exact"), ("1.Autocad", "AutoCAD", "exact"),
    (".NET", ".NET", "exact"), ("Statistics & Probability Basics", "Statistics", "alias"),
    ("Crop yield modelling", "Crop yield modelling", "unmapped"),
])
def test_skills_are_normalised_or_honestly_left_unmapped(raw, canonical, method):
    (skill,) = normalize_skills([raw])
    assert (skill.canonical, skill.method) == (canonical, method)


def test_compound_items_split_only_when_every_part_is_known():
    assert [s.canonical for s in normalize_skills(["Git/GitHub"])] == ["Git", "GitHub"]
    assert [s.canonical for s in normalize_skills(["CI/CD"])] == ["CI/CD"]
    assert [s.canonical for s in normalize_skills(["Python (Pandas, NumPy)"])][:2] == ["Python", "Pandas, NumPy"]


def test_stated_experience_is_kept_as_evidence():
    (skill,) = normalize_skills(["Python- Exprience - Less than 1 year months"])
    assert skill.canonical == "Python" and skill.evidence == "Less than 1 year"


# -- TF-IDF baseline -------------------------------------------------------------------------------
def test_tfidf_baseline_is_deterministic_and_reports_no_signal(embedder):
    from app.embeddings.provider import TfidfBaselineProvider

    again = TfidfBaselineProvider(embedder._dir)
    assert again.describe() == embedder.describe()  # same data -> same version and dimension
    vectors = embedder.embed(["python sql machine learning", "zzqx wvyk"])
    assert abs(sum(x * x for x in vectors[0].values) - 1) < 1e-5 and vectors[1] is None
    assert vectors[0].indices and len(vectors[0].indices) < 50  # sparse: a handful of non-zero terms (ADR 0011)


# -- uploaded resumes (2026-10-15): headings, entries, grouped skills --------------------------------------------
RESUME = """EDUCATION
Example Institute of Technology Mumbai, India
B.Tech, Computer Science & Engineering (Data Science) | CGPA: 8.5 / 10 Aug 2023 – Aug 2027
TECHNICAL SKILLS
Analytics & Reporting: Root Cause Analysis, Python (Pandas, NumPy), SQL,
Power BI (DAX, Star Schema), Excel (Advanced) – automated dashboards
PROJECTS
Procurement Risk System May 2026
• Built a risk scoring model across price, delivery and quality for 400+ purchase
orders; presented findings to the CFO.
• Designed an agentic pipeline via
LLM to classify supplier risk.
Field Force Analytics April 2026
• Applied K-Means clustering to classify representatives.
CERTIFICATIONS
Data Visualisation: Empowering Business – Tata iQ | Forage May 2026
LEADERSHIP EXPERIENCE
Vice Chairperson – NSS June 2025 – present
Example College Mumbai, India
• Coordinated a blood donation programme.
"""


def test_uploaded_resume_entries_are_one_item_each_with_wrapped_lines_joined():
    r = sections.resume_sections(RESUME)
    assert r["projects"] == [
        "Procurement Risk System May 2026: Built a risk scoring model across price, delivery and quality for 400+ "
        "purchase orders; presented findings to the CFO. Designed an agentic pipeline via LLM to classify supplier risk.",
        "Field Force Analytics April 2026: Applied K-Means clustering to classify representatives.",
    ]
    assert r["experience"] == ["Vice Chairperson – NSS June 2025 – present · Example College Mumbai, India: "
                               "Coordinated a blood donation programme."]
    assert r["certifications"] == ["Data Visualisation: Empowering Business – Tata iQ | Forage May 2026"]
    assert r["degree"] == "B.TECH" and "Computer Science" in r["education"]


def test_uploaded_resume_skills_drop_category_prefixes_and_expand_groups():
    assert sections.resume_skill_items(RESUME) == [
        "Root Cause Analysis", "Python", "Pandas", "NumPy", "SQL", "Power BI", "DAX", "Star Schema", "Excel",
        "automated dashboards"]
